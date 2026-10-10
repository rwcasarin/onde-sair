"""Testes do ranking com lugares fictícios: python3 -m unittest curadoria/test_ranking.py"""
import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ranking  # noqa: E402

CFG = {
    "cidade": "Sorocaba", "uf": "SP", "centro": {"lat": -23.5, "lng": -47.45}, "raio_m": 15000,
    "tipos": {
        "Restaurantes": {"cota": 2, "min_avaliacoes": 200, "buscas": ["restaurante"]},
        "Bares": {"cota": 1, "min_avaliacoes": 200, "buscas": ["bar"]},
    },
    "nota_minima": 4.3,
    "apostas": {"nota_minima": 4.6, "min_avaliacoes": 50, "maximo": 1},
    "ja_no_site": ["Frater's Pub"],
}


def lugar(pid, nome, nota, n, tipo="restaurant", cidade="Sorocaba", status="OPERATIONAL", bairro="Centro"):
    return {
        "id": pid, "displayName": {"text": nome}, "rating": nota, "userRatingCount": n,
        "businessStatus": status, "primaryType": tipo, "priceLevel": "PRICE_LEVEL_MODERATE",
        "googleMapsUri": f"https://maps.google.com/?cid={pid}", "formattedAddress": f"Rua X, 1 - {bairro}",
        "addressComponents": [
            {"longText": bairro, "types": ["sublocality_level_1", "sublocality"]},
            {"longText": cidade, "types": ["administrative_area_level_2", "political"]},
        ],
    }


def cenario():
    restaurantes = [
        lugar("r1", "Muito avaliado", 4.7, 3000),
        lugar("r2", "Poucas e perfeitas", 4.9, 210),
        lugar("r3", "Bom e popular", 4.4, 1500),
        lugar("r4", "Nota baixa", 4.1, 5000),
        lugar("r5", "Aposta forte", 4.8, 90),
        lugar("r6", "Aposta fraca", 4.7, 60),
        lugar("r7", "Fechou", 4.9, 800, status="CLOSED_PERMANENTLY"),
        lugar("r8", "Votorantim", 4.9, 900, cidade="Votorantim"),
        lugar("b9", "Bar que veio na busca de restaurante", 4.7, 700, tipo="bar"),
    ]
    bares = [
        lugar("b1", "Frater's Pub", 4.9, 900, tipo="pub"),
        lugar("b2", "Boteco", 4.4, 400, tipo="bar"),
        lugar("b9", "Bar que veio na busca de restaurante", 4.7, 700, tipo="bar"),
    ]
    return [("Restaurantes", "restaurante", restaurantes), ("Bares", "bar", bares)]


class TestRanking(unittest.TestCase):
    def setUp(self):
        self.linhas = ranking.consolidar(cenario(), CFG)
        self.stats = ranking.ranquear(self.linhas, CFG)
        self.por = {l["id"]: l for l in self.linhas}

    def test_ponderada_puxa_poucas_avaliacoes_para_a_media(self):
        self.assertAlmostEqual(ranking.ponderada(5.0, 0, 100, 4.0), 4.0)
        self.assertAlmostEqual(ranking.ponderada(4.0, 100, 100, 5.0), 4.5)
        self.assertEqual(ranking.ponderada(4.2, 0, 0, 0), 4.2)
        # nota um pouco menor com muitas avaliações vence nota alta com poucas
        self.assertGreater(ranking.ponderada(4.7, 3000, 855, 4.6), ranking.ponderada(4.9, 210, 855, 4.6))
        self.assertGreater(self.por["r1"]["pontuacao"], self.por["r2"]["pontuacao"])

    def test_cortes_cotas_e_reservas(self):
        sel = [l["id"] for l in self.linhas if l["tipo"] == "Restaurantes" and l["situacao"] == "Selecionado"]
        self.assertEqual(sorted(sel), ["r1", "r2"])
        self.assertEqual(self.por["r3"]["situacao"], "Reserva")
        self.assertEqual(self.por["r3"]["posicao"], 3)
        self.assertEqual(self.stats["Restaurantes"]["selecionados"], 2)
        self.assertEqual(self.stats["Restaurantes"]["reservas"], 1)
        self.assertIn("Abaixo do corte", self.por["r4"]["motivo"])

    def test_apostas_respeitam_limite(self):
        self.assertEqual(self.por["r5"]["situacao"], "Aposta")
        self.assertEqual(self.por["r6"]["situacao"], "Fora")
        self.assertEqual(self.por["r6"]["motivo"], "Aposta além do limite")

    def test_descartes(self):
        self.assertEqual(self.por["r7"]["motivo"], "Fechado")
        self.assertTrue(self.por["r8"]["motivo"].startswith("Outra cidade"))
        self.assertEqual(self.por["b1"]["motivo"], "Já está no site")

    def test_duplicata_vira_um_registro_com_tipo_do_google(self):
        b9 = [l for l in self.linhas if l["id"] == "b9"]
        self.assertEqual(len(b9), 1)
        self.assertEqual(b9[0]["tipo"], "Bares")
        self.assertEqual(b9[0]["buscas"], "restaurante, bar")
        self.assertEqual(b9[0]["situacao"], "Selecionado")   # vence o Boteco pela pontuação
        self.assertEqual(self.por["b2"]["situacao"], "Reserva")

    def test_tipo_pelo_google(self):
        self.assertEqual(ranking.tipo_pelo_google("japanese_restaurant"), "Restaurantes")
        self.assertEqual(ranking.tipo_pelo_google("night_club"), "Shows e baladas")
        self.assertEqual(ranking.tipo_pelo_google("park"), "Parques")
        self.assertIsNone(ranking.tipo_pelo_google("store"))

    def test_planilha_e_linha_de_comando(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            dados = Path(d, "dados.json")
            dados.write_text(json.dumps([{"tipo": t, "consulta": q, "lugares": ls} for t, q, ls in cenario()]))
            cidade = ranking.AQUI / "cidades" / "_teste.json"
            cidade.write_text(json.dumps(CFG))
            try:
                saida = Path(d, "r.xlsx")
                self.assertEqual(ranking.main(["_teste", "--dados", str(dados), "--saida", str(saida)]), 0)
            finally:
                cidade.unlink()
            wb = load_workbook(saida)
            self.assertEqual(wb.sheetnames, ["Resumo", "Ranking", "Descartados"])
            ws = wb["Ranking"]
            self.assertEqual(ws["A1"].value, "Decisão")
            nomes = [ws.cell(r, 5).value for r in range(2, ws.max_row + 1)]
            situacoes = [ws.cell(r, 4).value for r in range(2, ws.max_row + 1)]
            self.assertEqual(situacoes[-1], "Aposta")                       # apostas no fim
            self.assertEqual(nomes[:2], ["Muito avaliado", "Poucas e perfeitas"])  # ordem por pontuação
            self.assertTrue(ws.data_validations.dataValidation)
            self.assertEqual(wb["Descartados"].max_row - 1, 5)


if __name__ == "__main__":
    unittest.main()
