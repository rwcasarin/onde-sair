#!/usr/bin/env python3
"""
Onde Sair · ranking por cidade a partir de dados públicos da web (sem API)

Lê a coleta manual em curadoria/web/<cidade>.py (notas do Tripadvisor, do Google
via Wanderlog/Restaurant Guru e do Foursquare) e gera a planilha de curadoria.

    python3 curadoria/ranking_web.py sorocaba

Como as fontes têm escalas de volume muito diferentes (o Tripadvisor tem uma
fração das avaliações do Google), a média ponderada é calculada por fonte e as
fontes são combinadas com peso pela confiança de cada uma: v / (v + m).
"""
import json
import statistics
import sys
from datetime import date
from pathlib import Path

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
from ranking import ponderada, norm  # noqa: E402

FONTES = {"ta": "Tripadvisor", "g": "Google", "rg": "Restaurant Guru", "fq": "Foursquare"}
# volume mínimo de avaliações por fonte para entrar no ranking principal
MIN_AVAL = {"ta": 50, "g": 200, "rg": 200, "fq": 100}
# apostas: nota alta com volume entre este piso e o mínimo da fonte
PISO_APOSTA = {"ta": 15, "g": 50, "rg": 50, "fq": 30}

COTAS = {"Restaurantes": 35, "Bares": 25, "Parques": 12, "Shows e baladas": 10, "Eventos": 10}
# Atrações (parques, museus, teatros) têm notas mais baixas no Tripadvisor que restaurantes e bares
NOTA_MIN = {"Restaurantes": 4.3, "Bares": 4.3, "Shows e baladas": 4.3, "Parques": 4.0, "Eventos": 4.0}
APOSTA_NOTA, APOSTA_MAX = 4.6, 10
SELECIONADO, RESERVA, APOSTA, SEM_NOTA = "Selecionado", "Reserva", "Aposta", "Sem nota"


def carregar(cidade):
    ns = {}
    exec((AQUI / "web" / f"{cidade}.py").read_text(encoding="utf-8"), ns)
    return ns["L"]


def avaliar(lugares, ja_no_site=()):
    no_site = {norm(n) for n in ja_no_site}
    # parâmetros por fonte: nota média (C) e mediana de avaliações (m) entre os candidatos
    params = {}
    for f in FONTES:
        pares = [x["fontes"][f] for x in lugares if f in x["fontes"] and x["fontes"][f][0] and x["fontes"][f][1]]
        if pares:
            params[f] = {"C": statistics.mean(n for n, _ in pares), "m": statistics.median(v for _, v in pares)}

    for x in lugares:
        x.update(situacao="Fora", motivo="", pontuacao=None, posicao=None, aval_total=0)
        obs = (x.get("obs") or "").upper()
        if "FECHADO" in obs:
            x["motivo"] = "Fechado (segundo a fonte)"
            continue
        if "JÁ NO SITE" in obs or norm(x["nome"]) in no_site:
            x["motivo"] = "Já está no site"
            continue
        soma_w = soma = 0.0
        principal = aposta = False
        for f, (nota, v) in x["fontes"].items():
            if not nota:
                continue
            v = v or 0
            x["aval_total"] += v
            if f in params and v:
                p = params[f]
                w = v / (v + p["m"])
                soma += w * ponderada(nota, v, p["m"], p["C"])
                soma_w += w
            if v >= MIN_AVAL[f] and nota >= NOTA_MIN[x["tipo"]]:
                principal = True
            if nota >= APOSTA_NOTA and PISO_APOSTA[f] <= v < MIN_AVAL[f]:
                aposta = True
        if soma_w:
            x["pontuacao"] = round(soma / soma_w, 3)
        if not any(n for n, _ in x["fontes"].values() if n):
            x["situacao"] = SEM_NOTA
            x["motivo"] = "Nenhuma nota encontrada na busca"
        elif principal:
            x["situacao"] = "_aprovado"
        elif aposta:
            x["situacao"] = "_aposta"
        else:
            x["motivo"] = f"Abaixo do corte (nota {NOTA_MIN[x['tipo']]} com volume mínimo)"

    stats = {}
    for tipo, cota in COTAS.items():
        ap = sorted((x for x in lugares if x["tipo"] == tipo and x["situacao"] == "_aprovado"),
                    key=lambda x: (-(x["pontuacao"] or 0), -x["aval_total"]))
        for i, x in enumerate(ap, 1):
            x["posicao"], x["situacao"] = i, SELECIONADO if i <= cota else RESERVA
        stats[tipo] = {"cota": cota, "aprovados": len(ap), "selecionados": min(cota, len(ap)),
                       "falta": max(0, cota - len(ap)),
                       "sem_nota": sum(1 for x in lugares if x["tipo"] == tipo and x["situacao"] == SEM_NOTA)}
    apostas = sorted((x for x in lugares if x["situacao"] == "_aposta"),
                     key=lambda x: (-max(n for n, _ in x["fontes"].values() if n), -x["aval_total"]))
    for i, x in enumerate(apostas, 1):
        if i <= APOSTA_MAX:
            x["situacao"], x["posicao"] = APOSTA, i
        else:
            x["situacao"], x["motivo"] = "Fora", "Aposta além do limite"
    return stats, params


def salvar(lugares, stats, params, cidade, destino):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    fonte, negrito = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cab = PatternFill("solid", fgColor="4B2A7B")
    cor = {SELECIONADO: "E2F0D9", RESERVA: "F2F2F2", APOSTA: "FFF2CC", SEM_NOTA: "FCE4D6"}
    tipos = list(COTAS)

    wb = Workbook()
    rs = wb.active
    rs.title = "Resumo"
    linhas = [
        [f"Ranking de lugares · {cidade}"],
        [f"Gerado em {date.today():%d/%m/%Y} a partir de dados públicos da web (busca), sem API"],
        ["USO INTERNO: as notas são de terceiros (Tripadvisor, Google, Restaurant Guru) e não devem ser publicadas no site."],
        [],
        ["Como usar: na aba Ranking, preencha Decisão (Sim / Não / Talvez). Reservas substituem recusas. "
         "A aba Sem nota lista lugares citados em guias, mas sem nota encontrada: avalie pelo conhecimento local."],
        ["Corte: nota mínima por tipo em ao menos uma fonte, com volume mínimo nessa fonte "
         f"(Tripadvisor {MIN_AVAL['ta']}, Google {MIN_AVAL['g']}, Restaurant Guru {MIN_AVAL['rg']}, Foursquare {MIN_AVAL['fq']} avaliações)."],
        ["Pontuação: média ponderada (bayesiana) por fonte, combinada com peso pela confiança de cada fonte."],
        [f"Apostas: nota ≥ {APOSTA_NOTA} com volume abaixo do mínimo (piso: Tripadvisor {PISO_APOSTA['ta']}, Google {PISO_APOSTA['g']}); até {APOSTA_MAX}."],
        ["Limitação: dados vêm de resumos de busca, com datas diferentes por fonte. Confira nota e funcionamento antes de publicar."],
        [],
        ["Tipo", "Cota", "Nota mínima", "Aprovados", "Selecionados", "Faltam p/ cota", "Sem nota (avaliar)"],
    ]
    for t in tipos:
        s = stats[t]
        linhas.append([t, s["cota"], NOTA_MIN[t], s["aprovados"], s["selecionados"], s["falta"], s["sem_nota"]])
    linhas.append(["Apostas", APOSTA_MAX, APOSTA_NOTA, "", sum(1 for x in lugares if x["situacao"] == APOSTA)])
    linhas += [[], ["Fonte", "Nota média (C)", "Mediana de avaliações (m)", "Lugares com nota"]]
    for f, p in params.items():
        linhas.append([FONTES[f], round(p["C"], 2), p["m"], sum(1 for x in lugares if x["fontes"].get(f, (None,))[0])])
    for r in linhas:
        rs.append(r)
    for row in rs.iter_rows():
        for c in row:
            c.font = fonte
    rs["A1"].font = Font(name="Arial", size=13, bold=True)
    rs["A3"].font = Font(name="Arial", size=10, bold=True, color="C00000")
    for r in (11, 11 + len(tipos) + 3):
        for c in rs[r]:
            if c.value:
                c.font, c.fill = negrito, cab
    for col, w in zip("ABCDEFG", [24, 14, 24, 16, 13, 14, 18]):
        rs.column_dimensions[col].width = w

    def nv(x, f):
        n, v = x["fontes"].get(f, (None, None))
        return (n, v if v else None) if n else (None, None)

    colunas = [("Decisão", 10), ("Tipo", 15), ("Posição", 8), ("Situação", 12), ("Nome", 38), ("Subtipo", 24),
               ("Pontuação", 10), ("Tripadvisor nota", 11), ("Tripadvisor aval.", 11), ("Google nota", 10),
               ("Google aval.", 10), ("Rest. Guru nota", 10), ("Rest. Guru aval.", 10), ("Foursquare nota", 11),
               ("Foursquare aval.", 11), ("Observação / motivo", 50)]

    def aba(nome, itens, decisao=True):
        ws = wb.create_sheet(nome)
        ws.append([c for c, _ in colunas])
        for c in ws[1]:
            c.font, c.fill = negrito, cab
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for x in itens:
            vals = ["", x["tipo"], x["posicao"], x["situacao"], x["nome"], x["sub"], x["pontuacao"]]
            for f in ("ta", "g", "rg", "fq"):
                vals += list(nv(x, f))
            vals.append(" · ".join(filter(None, [x.get("obs"), x.get("motivo")])))
            ws.append(vals)
            r = ws.max_row
            for c in ws[r]:
                c.font = fonte
                c.alignment = Alignment(vertical="top")
            if x["situacao"] in cor:
                ws.cell(r, 4).fill = PatternFill("solid", fgColor=cor[x["situacao"]])
        for i, (_, w) in enumerate(colunas, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "F2"
        ws.auto_filter.ref = ws.dimensions
        if decisao and ws.max_row > 1:
            dv = DataValidation(type="list", formula1='"Sim,Não,Talvez"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"A2:A{ws.max_row}")
            for r in range(2, ws.max_row + 1):
                ws.cell(r, 1).fill = PatternFill("solid", fgColor="FFF2CC")

    ordem = [SELECIONADO, RESERVA, APOSTA]
    aba("Ranking", sorted((x for x in lugares if x["situacao"] in ordem),
                          key=lambda x: (x["situacao"] == APOSTA, tipos.index(x["tipo"]), ordem.index(x["situacao"]), x["posicao"])))
    aba("Sem nota", sorted((x for x in lugares if x["situacao"] == SEM_NOTA), key=lambda x: (tipos.index(x["tipo"]), x["nome"])))
    aba("Descartados", sorted((x for x in lugares if x["situacao"] == "Fora"),
                              key=lambda x: (tipos.index(x["tipo"]), x["motivo"], -(x["pontuacao"] or 0))), decisao=False)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if not argv:
        raise SystemExit("uso: ranking_web.py <cidade> [saida.xlsx]")
    cidade = argv[0]
    lugares = carregar(cidade)
    cfg_path = AQUI / "cidades" / f"{cidade}.json"
    cfg = json.loads(cfg_path.read_text(encoding="utf-8")) if cfg_path.exists() else {"cidade": cidade.title()}
    stats, params = avaliar(lugares, cfg.get("ja_no_site", []))
    destino = Path(argv[1]) if len(argv) > 1 else AQUI / "saida" / f"ranking-{cidade}-web.xlsx"
    salvar(lugares, stats, params, cfg.get("cidade", cidade.title()), destino)
    for t, s in stats.items():
        print(f"{t:16} cota {s['cota']:3} · aprovados {s['aprovados']:3} · faltam {s['falta']:2} · sem nota {s['sem_nota']}")
    print(f"apostas: {sum(1 for x in lugares if x['situacao'] == APOSTA)} → {destino}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
