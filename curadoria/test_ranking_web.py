"""Testes do ranking com dados da web: python3 -m unittest curadoria/test_ranking_web.py"""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ranking_web as rw  # noqa: E402


def lugar(nome, tipo, fontes, obs=""):
    return {"nome": nome, "tipo": tipo, "sub": "", "fontes": fontes, "obs": obs}


def cenario():
    return [
        lugar("Muito avaliado no Google", "Restaurantes", {"g": (4.6, 3000)}),
        lugar("Bem avaliado no Tripadvisor", "Restaurantes", {"ta": (4.8, 400)}),
        lugar("Notas divergentes", "Restaurantes", {"g": (4.6, 3500), "ta": (3.9, 80)}),
        lugar("Pouco volume, nota alta", "Restaurantes", {"ta": (4.8, 30)}),
        lugar("Volume minúsculo", "Restaurantes", {"ta": (5.0, 3)}),
        lugar("Nota baixa", "Restaurantes", {"ta": (4.0, 900)}),
        lugar("Só citado em guia", "Bares", {}),
        lugar("Fechou", "Bares", {"g": (4.7, 900)}, obs="FECHADO definitivamente"),
        lugar("Parque com nota de atração", "Parques", {"ta": (4.1, 200)}),
        lugar("Já publicado", "Bares", {"g": (4.8, 900)}),
    ]


class TestRankingWeb(unittest.TestCase):
    def setUp(self):
        self.L = cenario()
        self.stats, self.params = rw.avaliar(self.L, ["Já publicado"])
        self.por = {x["nome"]: x for x in self.L}

    def test_situacoes(self):
        s = {n: x["situacao"] for n, x in self.por.items()}
        self.assertEqual(s["Muito avaliado no Google"], "Selecionado")
        self.assertEqual(s["Bem avaliado no Tripadvisor"], "Selecionado")
        self.assertEqual(s["Pouco volume, nota alta"], "Aposta")
        self.assertEqual(s["Volume minúsculo"], "Fora")          # abaixo do piso de aposta
        self.assertEqual(s["Nota baixa"], "Fora")
        self.assertEqual(s["Só citado em guia"], "Sem nota")
        self.assertEqual(s["Fechou"], "Fora")
        self.assertEqual(s["Já publicado"], "Fora")
        self.assertEqual(self.por["Já publicado"]["motivo"], "Já está no site")

    def test_parques_tem_corte_proprio(self):
        self.assertEqual(self.por["Parque com nota de atração"]["situacao"], "Selecionado")

    def test_fontes_divergentes_baixam_a_pontuacao(self):
        self.assertLess(self.por["Notas divergentes"]["pontuacao"], self.por["Muito avaliado no Google"]["pontuacao"])

    def test_cotas_e_faltas(self):
        self.assertEqual(self.stats["Restaurantes"]["aprovados"], 3)
        self.assertEqual(self.stats["Restaurantes"]["falta"], rw.COTAS["Restaurantes"] - 3)
        self.assertEqual(self.stats["Bares"]["sem_nota"], 1)

    def test_planilha(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            destino = Path(d, "x.xlsx")
            rw.salvar(self.L, self.stats, self.params, "Teste", destino)
            wb = load_workbook(destino)
            self.assertEqual(wb.sheetnames, ["Resumo", "Ranking", "Sem nota", "Descartados"])
            nomes = [wb["Ranking"].cell(r, 5).value for r in range(2, wb["Ranking"].max_row + 1)]
            self.assertEqual(nomes[-1], "Pouco volume, nota alta")       # apostas no fim
            self.assertIn("Só citado em guia", [wb["Sem nota"].cell(r, 5).value for r in range(2, wb["Sem nota"].max_row + 1)])


if __name__ == "__main__":
    unittest.main()
