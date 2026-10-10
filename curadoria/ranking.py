#!/usr/bin/env python3
"""
Onde Sair · ranking de lugares por cidade (Google Places API New)

Busca candidatos no Google por tipo de lugar, aplica os cortes e as cotas da
cidade (curadoria/cidades/<cidade>.json) e gera uma planilha para a curadoria
marcar Sim / Não / Talvez.

    GOOGLE_PLACES_API_KEY=... python3 curadoria/ranking.py sorocaba
    python3 curadoria/ranking.py sorocaba --simular      # só conta as consultas
    python3 curadoria/ranking.py sorocaba --dados x.json  # usa lugares já baixados (testes)

Uso interno: os termos do Google não permitem publicar as notas nem guardar o
conteúdo por mais de 30 dias. A planilha serve só para escolher os lugares; no
site entra apenas o placeId (e os nossos textos).
"""
import argparse
import hashlib
import json
import os
import statistics
import sys
import time
import unicodedata
import urllib.error
import urllib.request
from datetime import date
from pathlib import Path

AQUI = Path(__file__).resolve().parent
API = "https://places.googleapis.com/v1/places:searchText"
CAMPOS = ",".join([
    "places.id", "places.displayName", "places.formattedAddress", "places.addressComponents",
    "places.location", "places.rating", "places.userRatingCount", "places.businessStatus",
    "places.priceLevel", "places.primaryType", "places.primaryTypeDisplayName", "places.googleMapsUri",
    "nextPageToken",
])
PAGINAS = 3            # a API devolve no máximo 20 por página e 60 por busca
CACHE_DIAS = 30        # limite dos termos do Google para guardar resultados

PRECO = {"PRICE_LEVEL_FREE": 0, "PRICE_LEVEL_INEXPENSIVE": 1, "PRICE_LEVEL_MODERATE": 2,
         "PRICE_LEVEL_EXPENSIVE": 3, "PRICE_LEVEL_VERY_EXPENSIVE": 3}

# Situações, na ordem em que aparecem na planilha
SELECIONADO, RESERVA, APOSTA = "Selecionado", "Reserva", "Aposta"
ORDEM_SITUACAO = [SELECIONADO, RESERVA, APOSTA]


def norm(s=""):
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return " ".join("".join(c if c.isalnum() else " " for c in s.lower()).split())


def tipo_pelo_google(primary_type):
    """Tipo do site sugerido pelo tipo principal do Google (None quando não dá pra dizer)."""
    t = primary_type or ""
    if t in ("night_club", "karaoke", "dance_hall", "live_music_venue", "concert_hall", "event_venue"):
        return "Shows e baladas"
    if t in ("bar", "pub", "wine_bar", "cocktail_bar", "brewery", "brewpub", "beer_garden", "sports_bar", "lounge_bar"):
        return "Bares"
    if t in ("park", "national_park", "state_park", "city_park", "hiking_area", "plaza", "botanical_garden", "garden", "observation_deck"):
        return "Parques"
    if t in ("museum", "art_gallery", "performing_arts_theater", "cultural_center", "movie_theater", "zoo", "aquarium", "historical_landmark", "planetarium"):
        return "Eventos"
    if t.endswith("restaurant") or t in ("cafe", "coffee_shop", "bakery", "ice_cream_shop", "dessert_shop", "confectionery", "pizza_place", "steak_house", "sandwich_shop", "food_court"):
        return "Restaurantes"
    return None


# ---------------------------------------------------------------------------
# Coleta
# ---------------------------------------------------------------------------
def chamar_api(corpo, chave):
    req = urllib.request.Request(API, data=json.dumps(corpo).encode(), method="POST", headers={
        "content-type": "application/json", "X-Goog-Api-Key": chave, "X-Goog-FieldMask": CAMPOS,
    })
    for tentativa in range(4):
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                return json.load(r)
        except urllib.error.HTTPError as e:
            msg = e.read().decode(errors="replace")[:400]
            if e.code in (429, 500, 503) and tentativa < 3:
                time.sleep(2 ** (tentativa + 1))
                continue
            raise SystemExit(f"Erro {e.code} da Places API: {msg}")


def buscar(consulta, cfg, chave, cache_dir):
    """Todas as páginas de uma busca, com cache local de até 30 dias."""
    base = {
        "textQuery": f"{consulta} em {cfg['cidade']} {cfg['uf']}",
        "languageCode": "pt-BR", "regionCode": "BR", "pageSize": 20,
        "locationBias": {"circle": {"center": {"latitude": cfg["centro"]["lat"], "longitude": cfg["centro"]["lng"]},
                                    "radius": float(cfg["raio_m"])}},
    }
    arq = cache_dir / (hashlib.sha1(json.dumps(base, sort_keys=True).encode()).hexdigest()[:16] + ".json")
    if arq.exists() and time.time() - arq.stat().st_mtime < CACHE_DIAS * 86400:
        return json.loads(arq.read_text())
    lugares, token = [], None
    for _ in range(PAGINAS):
        corpo = dict(base, pageToken=token) if token else base
        r = chamar_api(corpo, chave)
        lugares += r.get("places", [])
        token = r.get("nextPageToken")
        if not token:
            break
        time.sleep(2)  # o token da próxima página leva alguns segundos para valer
    cache_dir.mkdir(parents=True, exist_ok=True)
    arq.write_text(json.dumps(lugares, ensure_ascii=False))
    return lugares


def coletar(cfg, chave, cache_dir):
    """[(tipo, consulta, [lugares])] para todas as buscas da cidade."""
    out = []
    for tipo, t in cfg["tipos"].items():
        for q in t["buscas"]:
            achados = buscar(q, cfg, chave, cache_dir)
            print(f"  {tipo:16} {q:28} {len(achados):3} lugares", file=sys.stderr)
            out.append((tipo, q, achados))
    return out


# ---------------------------------------------------------------------------
# Normalização e ranking
# ---------------------------------------------------------------------------
def componente(p, *tipos):
    for t in tipos:
        for c in p.get("addressComponents", []):
            if t in c.get("types", []):
                return c.get("longText") or c.get("shortText") or ""
    return ""


def consolidar(buscas, cfg):
    """Um registro por placeId. O tipo vem do Google quando ele é claro; senão, da 1ª busca que achou o lugar."""
    por_id = {}
    for tipo, q, lugares in buscas:
        for p in lugares:
            pid = p.get("id")
            if not pid:
                continue
            if pid not in por_id:
                por_id[pid] = {"p": p, "tipos_busca": [], "buscas": []}
            r = por_id[pid]
            if tipo not in r["tipos_busca"]:
                r["tipos_busca"].append(tipo)
            r["buscas"].append(q)
    linhas = []
    for pid, r in por_id.items():
        p = r["p"]
        sugerido = tipo_pelo_google(p.get("primaryType"))
        tipo = sugerido if sugerido in cfg["tipos"] else r["tipos_busca"][0]
        linhas.append({
            "id": pid,
            "nome": (p.get("displayName") or {}).get("text", ""),
            "tipo": tipo,
            "nota": p.get("rating"),
            "avaliacoes": p.get("userRatingCount") or 0,
            "status": p.get("businessStatus", ""),
            "preco": PRECO.get(p.get("priceLevel")),
            "bairro": componente(p, "sublocality_level_1", "sublocality", "neighborhood"),
            "cidade": componente(p, "administrative_area_level_2", "locality"),
            "endereco": p.get("formattedAddress", ""),
            "tipo_google": (p.get("primaryTypeDisplayName") or {}).get("text", "") or p.get("primaryType", ""),
            "maps": p.get("googleMapsUri", ""),
            "buscas": ", ".join(dict.fromkeys(r["buscas"])),
        })
    return linhas


def ponderada(nota, v, m, media):
    """Média bayesiana: com poucas avaliações, a nota puxa para a média do tipo na cidade."""
    if v + m == 0:
        return nota
    return v / (v + m) * nota + m / (v + m) * media


def ranquear(linhas, cfg):
    """Preenche situacao, motivo, pontuacao e posicao em cada linha. Devolve as estatísticas por tipo."""
    no_site = {norm(n) for n in cfg.get("ja_no_site", [])}
    cidade = norm(cfg["cidade"])
    for l in linhas:
        l.update(situacao="Fora", motivo="", pontuacao=None, posicao=None)
        if l["cidade"] and norm(l["cidade"]) != cidade:
            l["motivo"] = f"Outra cidade ({l['cidade']})"
        elif norm(l["nome"]) in no_site:
            l["motivo"] = "Já está no site"
        elif l["status"] and l["status"] != "OPERATIONAL":
            l["motivo"] = "Fechado" + (" temporariamente" if l["status"] == "CLOSED_TEMPORARILY" else "")
        elif l["nota"] is None:
            l["motivo"] = "Sem nota"

    stats, pool_apostas = {}, []
    ap = cfg["apostas"]
    for tipo, t in cfg["tipos"].items():
        validos = [l for l in linhas if l["tipo"] == tipo and not l["motivo"]]
        media = statistics.mean(l["nota"] for l in validos) if validos else 0
        m = statistics.median(l["avaliacoes"] for l in validos) if validos else 0
        stats[tipo] = {"candidatos": len(validos), "media": media, "m": m}
        for l in validos:
            l["pontuacao"] = round(ponderada(l["nota"], l["avaliacoes"], m, media), 3)
        aprovados = sorted((l for l in validos if l["nota"] >= cfg["nota_minima"] and l["avaliacoes"] >= t["min_avaliacoes"]),
                           key=lambda l: (-l["pontuacao"], -l["avaliacoes"]))
        for i, l in enumerate(aprovados, 1):
            l["posicao"] = i
            l["situacao"] = SELECIONADO if i <= t["cota"] else RESERVA
        for l in validos:
            if l["situacao"] != "Fora":
                continue
            if l["nota"] >= ap["nota_minima"] and ap["min_avaliacoes"] <= l["avaliacoes"] < t["min_avaliacoes"]:
                pool_apostas.append(l)
            else:
                l["motivo"] = f"Abaixo do corte (nota {cfg['nota_minima']} e {t['min_avaliacoes']} avaliações)"
        stats[tipo]["selecionados"] = sum(1 for l in aprovados if l["situacao"] == SELECIONADO)
        stats[tipo]["reservas"] = sum(1 for l in aprovados if l["situacao"] == RESERVA)

    pool_apostas.sort(key=lambda l: (-l["nota"], -l["avaliacoes"]))
    for i, l in enumerate(pool_apostas, 1):
        if i <= ap["maximo"]:
            l["situacao"], l["posicao"] = APOSTA, i
        else:
            l["motivo"] = "Aposta além do limite"
    return stats


# ---------------------------------------------------------------------------
# Planilha
# ---------------------------------------------------------------------------
def salvar(linhas, stats, cfg, destino, consultas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    fonte, negrito = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cab = PatternFill("solid", fgColor="4B2A7B")
    cor = {SELECIONADO: "E2F0D9", RESERVA: "F2F2F2", APOSTA: "FFF2CC"}
    tipos = list(cfg["tipos"])

    wb = Workbook()
    rs = wb.active
    rs.title = "Resumo"
    resumo = [
        [f"Ranking de lugares · {cfg['cidade']} ({cfg['uf']})"],
        [f"Gerado em {date.today():%d/%m/%Y} a partir da Google Places API · {consultas} buscas"],
        ["USO INTERNO: as notas do Google não podem ser publicadas no site nem guardadas por mais de 30 dias."],
        [],
        ["Como usar: na aba Ranking, preencha a coluna Decisão (Sim / Não / Talvez). Reservas substituem selecionados recusados."],
        [f"Corte: nota ≥ {cfg['nota_minima']} e avaliações mínimas por tipo. Pontuação = média ponderada (bayesiana) por tipo."],
        [f"Apostas: nota ≥ {cfg['apostas']['nota_minima']}, de {cfg['apostas']['min_avaliacoes']} avaliações até o mínimo do tipo; até {cfg['apostas']['maximo']} lugares."],
        [],
        ["Tipo", "Cota", "Mín. avaliações", "Candidatos", "Selecionados", "Reservas", "Nota média (C)", "Mediana de avaliações (m)"],
    ]
    for t in tipos:
        s, c = stats[t], cfg["tipos"][t]
        resumo.append([t, c["cota"], c["min_avaliacoes"], s["candidatos"], s["selecionados"], s["reservas"],
                       round(s["media"], 2), s["m"]])
    resumo.append(["Apostas", cfg["apostas"]["maximo"], "", "", sum(1 for l in linhas if l["situacao"] == APOSTA)])
    for r in resumo:
        rs.append(r)
    for row in rs.iter_rows():
        for c in row:
            c.font = fonte
    rs["A1"].font = Font(name="Arial", size=13, bold=True)
    rs["A3"].font = Font(name="Arial", size=10, bold=True, color="C00000")
    for c in rs[9]:
        c.font, c.fill = negrito, cab
    for col, w in zip("ABCDEFGH", [20, 8, 15, 12, 13, 10, 15, 24]):
        rs.column_dimensions[col].width = w

    colunas = [("Decisão", 11), ("Tipo", 16), ("Posição", 8), ("Situação", 12), ("Nome", 34), ("Nota", 6),
               ("Avaliações", 11), ("Pontuação", 10), ("Preço (0–3)", 10), ("Bairro", 20), ("Endereço", 44),
               ("Tipo no Google", 20), ("Google Maps", 16), ("Place ID", 30), ("Encontrado nas buscas", 36)]

    def aba(nome, itens, com_motivo=False):
        ws = wb.create_sheet(nome)
        cols = colunas + ([("Motivo", 40)] if com_motivo else [])
        ws.append([c for c, _ in cols])
        for c in ws[1]:
            c.font, c.fill = negrito, cab
        for l in itens:
            ws.append(["", l["tipo"], l["posicao"], l["situacao"], l["nome"], l["nota"], l["avaliacoes"], l["pontuacao"],
                       l["preco"], l["bairro"], l["endereco"], l["tipo_google"], "abrir" if l["maps"] else "", l["id"], l["buscas"]]
                      + ([l["motivo"]] if com_motivo else []))
            r = ws.max_row
            for c in ws[r]:
                c.font = fonte
            if l["situacao"] in cor:
                ws.cell(r, 4).fill = PatternFill("solid", fgColor=cor[l["situacao"]])
            if l["maps"]:
                ws.cell(r, 13).hyperlink = l["maps"]
                ws.cell(r, 13).font = Font(name="Arial", size=10, color="0563C1", underline="single")
        for i, (_, w) in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "F2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top")
        return ws

    escolhidos = sorted((l for l in linhas if l["situacao"] in ORDEM_SITUACAO),
                        key=lambda l: (ORDEM_SITUACAO.index(l["situacao"]) == 2, tipos.index(l["tipo"]),
                                       ORDEM_SITUACAO.index(l["situacao"]), l["posicao"]))
    ws = aba("Ranking", escolhidos)
    dv = DataValidation(type="list", formula1='"Sim,Não,Talvez"', allow_blank=True)
    ws.add_data_validation(dv)
    if ws.max_row > 1:
        dv.add(f"A2:A{ws.max_row}")
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 1).fill = PatternFill("solid", fgColor="FFF2CC")
    fora = sorted((l for l in linhas if l["situacao"] == "Fora"), key=lambda l: (tipos.index(l["tipo"]), l["motivo"], -(l["nota"] or 0)))
    aba("Descartados", fora, com_motivo=True)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


# ---------------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description="Ranking de lugares por cidade (Google Places API).")
    ap.add_argument("cidade", help="nome do arquivo em curadoria/cidades/, sem .json (ex.: sorocaba)")
    ap.add_argument("--saida", help="planilha de saída (padrão: curadoria/saida/ranking-<cidade>.xlsx)")
    ap.add_argument("--simular", action="store_true", help="só mostra quantas buscas seriam feitas")
    ap.add_argument("--dados", help="JSON com [{tipo, consulta, lugares}] em vez de chamar a API")
    a = ap.parse_args(argv)

    cfg = json.loads((AQUI / "cidades" / f"{a.cidade}.json").read_text(encoding="utf-8"))
    consultas = sum(len(t["buscas"]) for t in cfg["tipos"].values())
    if a.simular:
        print(f"{consultas} buscas · até {consultas * PAGINAS} chamadas à API (3 páginas por busca, no máximo)")
        return 0
    if a.dados:
        buscas = [(b["tipo"], b["consulta"], b["lugares"]) for b in json.loads(Path(a.dados).read_text(encoding="utf-8"))]
    else:
        chave = os.environ.get("GOOGLE_PLACES_API_KEY")
        if not chave:
            raise SystemExit("Defina GOOGLE_PLACES_API_KEY (chave da Places API New) no ambiente.")
        buscas = coletar(cfg, chave, AQUI / ".cache" / a.cidade)

    linhas = consolidar(buscas, cfg)
    stats = ranquear(linhas, cfg)
    destino = Path(a.saida) if a.saida else AQUI / "saida" / f"ranking-{a.cidade}.xlsx"
    salvar(linhas, stats, cfg, destino, consultas)
    sel = sum(1 for l in linhas if l["situacao"] == SELECIONADO)
    res = sum(1 for l in linhas if l["situacao"] == RESERVA)
    apo = sum(1 for l in linhas if l["situacao"] == APOSTA)
    print(f"{len(linhas)} lugares únicos · {sel} selecionados · {res} reservas · {apo} apostas → {destino}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
